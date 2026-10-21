"""Versioned original reference notes and user-selected document retrieval."""
from pathlib import Path
import hashlib,time,zipfile
from knowledge import KnowledgeStore,TextHTML,read_web

CATALOG=[
 ('Programación Python','https://docs.python.org/3/','Documentación oficial y biblioteca estándar'),
 ('Web y JavaScript','https://developer.mozilla.org/es/','Documentación de Mozilla'),
 ('Ciencias y matemáticas','https://openstax.org/subjects','Libros abiertos de Rice University'),
 ('Astronomía','https://science.nasa.gov/','NASA: misiones y ciencia'),
 ('Tierra y geología','https://www.usgs.gov/','USGS: datos y publicaciones'),
 ('Literatura en español','https://www.cervantesvirtual.com/','Biblioteca Virtual Miguel de Cervantes'),
 ('Libros de dominio público','https://www.gutenberg.org/','Project Gutenberg; disponibilidad según país'),
 ('Música e historia cultural','https://www.loc.gov/collections/','Colecciones de Library of Congress'),
 ('Historia y cultura','https://www.si.edu/','Colecciones del Smithsonian'),
 ('Investigación biomédica','https://pubmed.ncbi.nlm.nih.gov/','Índice de investigación; información sanitaria requiere verificación actual'),
 ('Economía y series','https://fred.stlouisfed.org/','Federal Reserve Bank of St. Louis'),
 ('Oro','https://www.gold.org/goldhub','World Gold Council: contexto y metodología'),
]
NOTES=[
 ('Matemáticas: riesgo y probabilidad','La esperanza de un resultado discreto es la suma de cada resultado multiplicado por su probabilidad. La varianza mide dispersión alrededor de la media. Una proporción observada necesita denominador e intervalo de incertidumbre. Correlación no establece causalidad. Un resultado positivo en una muestra usada para ajustar un sistema no demuestra rendimiento fuera de muestra.'),
 ('Programación: diseño y pruebas','Las funciones separan entrada, transformación y salida. Una prueba útil comprueba comportamiento observable, incluidos errores y límites. El estado compartido entre hilos requiere coordinación; una interfaz debe recibir resultados por mensajes y actualizar widgets en su propio hilo. Las consultas SQL usan parámetros y las escrituras importantes usan reemplazo atómico.'),
 ('Física: magnitudes y conservación','Una magnitud necesita número y unidad. La velocidad expresa cambio de posición por tiempo; la aceleración expresa cambio de velocidad. En un sistema aislado se conserva la energía total, aunque cambie de forma. Un modelo físico debe indicar supuestos, escala de aplicación e incertidumbre de medida.'),
 ('Música: ritmo, melodía y armonía','El pulso es la referencia temporal regular. El ritmo organiza duraciones y acentos. La melodía es una secuencia de alturas; la armonía estudia sus combinaciones. Una octava duplica la frecuencia. El timbre distingue instrumentos que tocan una misma altura, por su espectro y evolución temporal.'),
 ('Literatura: lectura y análisis','Conviene distinguir autor, narrador y personajes. La interpretación se apoya en el texto, sus recursos y contexto. Una metáfora establece una relación figurada; una comparación explicita esa relación. Una lectura crítica separa evidencia textual, interpretación propia y afirmaciones históricas que necesitan fuente.'),
 ('Ciencia: evidencia y actualización','Una hipótesis propone una explicación comprobable. Un experimento define método, variables y criterios antes de observar resultados. Replicación, transparencia y resultados contradictorios importan. Una noticia sobre un estudio no sustituye su diseño, muestra, limitaciones y publicación. Conocimiento reciente se verifica con fuentes fechadas.'),
 ('Trading: órdenes y ejecución','BUY LIMIT busca una compra por debajo del ask; BUY STOP una ruptura por encima. SELL LIMIT busca una venta por encima del bid; SELL STOP una ruptura por debajo. La entrada es una condición futura; el SL define invalidación. El riesgo monetario depende de contrato, lote y distancia, se calcula con MT5. Una estrella técnica no es una probabilidad calibrada. BTC y oro comparten capacidad de cuenta aunque sean símbolos distintos.'),
 ('Historia y cultura: fuentes','Una fuente primaria corresponde al periodo, documento o acontecimiento que se investiga. Una fuente secundaria interpreta otras fuentes. Autoría, fecha, procedencia y propósito ayudan a valorar sesgos. No deben confundirse una reconstrucción artística o ficción con evidencia histórica.'),
]

def seed_library(root):
 store=KnowledgeStore(root)
 for index,(title,body) in enumerate(NOTES):
  url='akira://biblioteca/16/'+str(index)
  if not store.get(url):store.put({'url':url,'title':title,'topic':title,'body':body,'summary':body,'kind':'reference','status':'indexed','published':'Guía original incluida; conceptos generales, no hechos actuales'})
 return store

def document_text(path):
 path=Path(path)
 if path.stat().st_size>20_000_000:raise ValueError('El documento supera 20 MB; divide el archivo antes de importarlo')
 ext=path.suffix.lower()
 if ext in ('.docx','.epub'):
  with zipfile.ZipFile(path) as archive:
   if sum(r.file_size for r in archive.infolist())>40_000_000:raise ValueError('Documento comprimido demasiado grande')
   if ext=='.epub':
    parser=TextHTML()
    for name in archive.namelist():
     if name.lower().endswith(('.html','.xhtml','.htm')):parser.feed(archive.read(name).decode('utf-8',errors='replace'))
    return '\n'.join(parser.parts)[:240000]
  from docx import Document
  doc=Document(path);return '\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows])[:240000]
 if ext=='.pdf':
  from pypdf import PdfReader
  reader=PdfReader(path)
  if reader.is_encrypted:raise ValueError('El PDF está cifrado; importa una copia que puedas leer')
  body='\n'.join(p.extract_text() or '' for p in reader.pages[:100])[:240000]
  if not body.strip():raise ValueError('PDF sin texto extraíble. Importa un PDF con OCR o texto')
  return body
 if ext not in ('.txt','.md','.csv','.py','.js','.html','.json'):raise ValueError('Formato no compatible con la biblioteca')
 return path.read_text(encoding='utf-8-sig',errors='replace')[:240000]

def import_document(root,path):
 path=Path(path);body=document_text(path)
 if not body.strip():raise ValueError('El documento está vacío')
 digest=hashlib.sha256(body.encode()).hexdigest();store=KnowledgeStore(root);count=0
 for offset in range(0,len(body),4000):
  chunk=body[offset:offset+4000];store.put({'url':'akira://documento/'+digest+'/'+str(offset),
   'title':path.name+' · fragmento '+str(offset//4000+1),'topic':'Documento importado por el usuario','body':chunk,
   'summary':chunk,'kind':'document','status':'indexed','published':'Sin fecha de publicación verificada'});count+=1
 return {'name':path.name,'chunks':count,'characters':len(body),'scope':'Hasta 100 páginas / 240.000 caracteres por archivo; los fragmentos conservan origen. No es entrenamiento del modelo.'}

def update_reference(root,url):
 if url not in [row[1] for row in CATALOG]:raise ValueError('Referencia fuera del catálogo')
 row=read_web(url)
 # A redirect must not turn the canonical catalog URL into a permanent cache
 # miss. Retain both requested and actually read URLs as provenance.
 row['resolved_url']=row['url'];row['url']=url
 row['topic']=next(r[0] for r in CATALOG if r[1]==url);row['kind']='reference_web';KnowledgeStore(root).put(row)
 return {'title':row['title'],'characters':len(row['body']),'fetched':row['fetched'],'scope':'Página pública consultada; no se afirma haber descargado todo el sitio'}
